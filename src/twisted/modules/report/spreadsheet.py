"""Findings + assets exporter to XLSX and CSV."""

from __future__ import annotations

import csv

from openpyxl import Workbook

from ...core.paths import to_canonical
from .._engine_callback import get_client
from ..base import EvidenceRef, ModuleContext, ModuleResult


def _fetch_engagement(client, engagement_id: int) -> dict:
    return {
        "engagement": client.get(f"/engagements/{engagement_id}"),
        "assets": client.get(f"/engagements/{engagement_id}/assets/detail"),
        "findings": client.get(f"/engagements/{engagement_id}/findings"),
    }


def _write_xlsx(path, payload: dict) -> None:
    wb = Workbook()
    # Findings sheet
    ws = wb.active
    ws.title = "Findings"
    ws.append(["id", "title", "severity", "cvss_score", "cvss_vector", "cwe", "status"])
    for f in payload["findings"]:
        ws.append([f["id"], f["title"], f["severity"], f.get("cvss_score"),
                   f.get("cvss_vector"), f.get("cwe"), f["status"]])

    # Assets sheet
    ws = wb.create_sheet("Assets")
    ws.append(["id", "host", "ip", "env_type", "in_scope", "risk_total",
               "open_ports", "techs", "cve_count"])
    for a in payload["assets"]:
        ws.append([
            a["id"], a["host"], a.get("ip"), a.get("env_type"),
            "yes" if a["in_scope"] else "no", a["risk_total"],
            ",".join(str(p["port"]) for p in a.get("ports", [])),
            ",".join(f"{t['name']}/{t.get('version','')}" for t in a.get("techs", [])),
            len(a.get("cves", [])),
        ])

    # CVEs sheet
    ws = wb.create_sheet("CVEs")
    ws.append(["asset_host", "cve_id", "cvss_score", "severity", "summary"])
    for a in payload["assets"]:
        for c in a.get("cves", []):
            ws.append([a["host"], c["cve_id"], c.get("cvss_score"),
                       c.get("severity"), c.get("summary")])

    wb.save(str(path))


def _write_csv(path, payload: dict) -> None:
    with open(path, "w", newline="") as fp:
        w = csv.writer(fp)
        w.writerow(["id", "title", "severity", "cvss_score", "cvss_vector", "cwe", "status"])
        for f in payload["findings"]:
            w.writerow([f["id"], f["title"], f["severity"], f.get("cvss_score"),
                        f.get("cvss_vector"), f.get("cwe"), f["status"]])


def export(ctx: ModuleContext) -> ModuleResult:
    engagement_id = (ctx.params or {}).get("engagement_id") or ctx.engagement_id
    if not engagement_id:
        return ModuleResult(success=False, error="missing required param 'engagement_id'")
    client = get_client()
    try:
        payload = _fetch_engagement(client, engagement_id)
    except Exception as e:  # noqa: BLE001
        return ModuleResult(success=False, error=f"engine fetch failed: {e}")

    ts = ctx.timestamp.strftime("%Y%m%d_%H%M%S")
    xlsx_path = ctx.work_dir / f"findings_{ts}.xlsx"
    csv_path = ctx.work_dir / f"findings_{ts}.csv"
    _write_xlsx(xlsx_path, payload)
    _write_csv(csv_path, payload)

    return ModuleResult(
        success=True,
        artifacts=[xlsx_path, csv_path],
        evidence=[
            EvidenceRef(path=to_canonical(xlsx_path), kind="file",
                        note="findings spreadsheet (xlsx)", host=ctx.worker_host),
            EvidenceRef(path=to_canonical(csv_path), kind="file",
                        note="findings spreadsheet (csv)", host=ctx.worker_host),
        ],
        summary=f"export: {len(payload['findings'])} finding(s), {len(payload['assets'])} asset(s)",
    )
